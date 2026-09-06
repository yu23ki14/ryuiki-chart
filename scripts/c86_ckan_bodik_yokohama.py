"""BODIK(神奈川県参加団体) と横浜市の CKAN 目録収穫、および実データ取得。

対象:
- ckan_bodik_kanagawa      : BODIK(全国共同CKAN https://data.bodik.jp) のうち神奈川県内の参加団体
                              (横須賀市142018 / 厚木市142123 / 真鶴町143839) の目録
- atsugi_river_water_quality: 厚木市「河川水質調査結果」(相模川ほか, XLSX, CC BY 4.0) を実際にダウンロードし
                              縦持ちに変換
- ckan_yokohama            : 横浜市オープンデータポータル (https://data.city.yokohama.lg.jp, 本物のCKAN) の目録
- yokohama_river_waterlevel: 横浜市「河川水位オープンデータ」(CSV, CC BY)
- yokohama_bio_indicator   : 横浜市「環境管理計画年次報告書 資料編」の「生物多様性」ZIP内、
                              河川の生物指標水質評価CSV

注意(地理的位置づけ): 横浜市データの対象河川(鶴見川・帷子川・大岡川・侍従川・境川・宮川)は
相模川水系ではない(境川も相模湾に注ぐ別水系)。流域カルテ本体(相模川水系)とは異なる隣接下流都市の
参考データとして扱うこと。register() の notes にもその旨を明記する。

c01_ckan.py / c02_ckan_env_select.py と同じ作法(common.py のみ経由でHTTP、出力はCSV+JSONL両方、
source_id/source_ref必須、register必須)に合わせる。既存ファイルは一切変更しない。
"""
import sys, re, json, pathlib, zipfile, datetime
sys.path.insert(0, str(pathlib.Path(__file__).parent))
from common import get, get_json, download, register, write_jsonl, \
                    to_fiscal_year, to_number, PROC, RAW, DB, appdb, now, sha256
import pandas as pd
import openpyxl

# ============================================================
# 0. 共通 CKAN 収穫ヘルパ (c01_ckan.py と同じロジック)
# ============================================================

def ckan_harvest(base, fq=None):
    out, start, rows = [], 0, 1000
    while True:
        params = {"rows": rows, "start": start}
        if fq: params["fq"] = fq
        j = get_json(f"{base}/api/3/action/package_search", params=params)
        res = j["result"]; out += res["results"]
        start += rows
        if start >= res["count"] or not res["results"]: break
    return out


def ckan_to_rows(base, sid, pkgs):
    """package_search の結果を ckan_datasets.csv / ckan_resources.csv と同じ列構成の行に変換"""
    ds_rows, res_rows = [], []
    for p in pkgs:
        ds_rows.append({
            "instance": sid, "dataset_id": p.get("id"), "name": p.get("name"),
            "title": p.get("title"), "notes": (p.get("notes") or "")[:800],
            "organization": (p.get("organization") or {}).get("title"),
            "license": p.get("license_title") or p.get("license_id"),
            "license_url": p.get("license_url"),
            "groups": "|".join(g.get("title", "") for g in p.get("groups") or []),
            "tags": "|".join(t.get("display_name", "") for t in p.get("tags") or []),
            "n_resources": len(p.get("resources") or []),
            "metadata_modified": p.get("metadata_modified"),
            "url": f"{base}/dataset/{p.get('name')}",
        })
        for r in p.get("resources") or []:
            res_rows.append({
                "instance": sid, "dataset_id": p.get("id"), "dataset_title": p.get("title"),
                "resource_id": r.get("id"), "resource_name": r.get("name"),
                "format": (r.get("format") or "").upper(), "url": r.get("url"),
                "size": r.get("size"), "last_modified": r.get("last_modified"),
                "license": p.get("license_title") or p.get("license_id"),
                "organization": (p.get("organization") or {}).get("title"),
                "groups": "|".join(g.get("title", "") for g in p.get("groups") or []),
            })
    return ds_rows, res_rows


# ============================================================
# 1. BODIK (神奈川県参加団体3団体) 目録収穫
# ============================================================

BODIK_BASE = "https://data.bodik.jp"
BODIK_ORGS = ["142018", "142123", "143839"]  # 横須賀市, 厚木市, 真鶴町
BODIK_ORG_NAMES = {"142018": "横須賀市", "142123": "厚木市", "143839": "真鶴町"}


def harvest_bodik():
    fq = "organization:(" + " OR ".join(BODIK_ORGS) + ")"
    pkgs = ckan_harvest(BODIK_BASE, fq=fq)
    ds_rows, res_rows = ckan_to_rows(BODIK_BASE, "bodik_kanagawa", pkgs)
    write_jsonl("ckan_datasets_bodik", ds_rows)
    write_jsonl("ckan_resources_bodik", res_rows)
    pd.DataFrame(ds_rows).to_csv(PROC / "ckan_datasets_bodik.csv", index=False)
    pd.DataFrame(res_rows).to_csv(PROC / "ckan_resources_bodik.csv", index=False)
    by_org = {}
    for p in pkgs:
        by_org[(p.get("organization") or {}).get("name")] = by_org.get((p.get("organization") or {}).get("name"), 0) + 1
    print(f"  bodik_kanagawa: {len(pkgs)} datasets  (by org: {by_org})")
    register("ckan_bodik_kanagawa", "BODIK オープンデータカタログ(神奈川県内参加団体)",
              "BODIK(自治体クラウド推進機構) / 横須賀市・厚木市・真鶴町",
              f"{BODIK_BASE}/organization/{'+'.join(BODIK_ORGS)}",
              "オープンデータ目録", "CKAN API (package_search, fq=organization絞込み)", "JSON",
              "各データセット個別(多くは CC BY 4.0)",
              True, len(pkgs),
              "c01_ckan.py に登録済みの base URL (https://odcs.bodik.jp/kanagawa) は誤り"
              "(404, WordPress製ランディングページでCKANでない)。正しい API base は https://data.bodik.jp。"
              "全国共同CKAN全18,390件のうち神奈川県関連の参加団体は3団体のみ"
              "(横須賀市142018/厚木市142123/真鶴町143839)。"
              "BODIKは全国共同CKANのため、参加団体名を fq=organization:(142018 OR 142123 OR 143839) で"
              "3団体に絞って収穫した(絞らないと18,390件全国分が混入する)。")
    return pkgs, ds_rows, res_rows


# ============================================================
# 2. 厚木市 河川水質調査結果 (atsugi_river_water_quality)
# ============================================================

ATSUGI_VARS = [
    # (openpyxl列インデックス(0始まり, iter_rows values_only の row タプル内), variable_ja, unit_ja)
    (2, "pH", ""),
    (3, "生物化学的酸素要求量 BOD", "mg/L"),
    (4, "化学的酸素要求量 COD", "mg/L"),
    (5, "浮遊物質量 SS", "mg/L"),
    (6, "溶存酸素量 DO", "mg/L"),
]
MONTH_HDR_RE = re.compile(r"((?:令和|平成|昭和|R|H|S)\s*(?:元|\d+)年)?\s*(\d+)月")
BELOW_DL_RE = re.compile(r"^([\d.]+)未満$")


def parse_atsugi_xlsx(path, resource_url, resource_name):
    """厚木市 河川水質調査結果 XLSX(月毎シート)を縦持ち行のリストに変換する。
    戻り値: (rows, issues) issues は統合できなかった/想定外だった箇所の説明文リスト。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet0 = wb.sheetnames[0]
    issues = []
    if not sheet0.startswith("月毎"):
        return [], [f"{resource_name}: 想定外のシート構成(先頭シートが「月毎」で始まらない): {wb.sheetnames}"]
    ws = wb[sheet0]
    title = ws.cell(1, 2).value or ""
    fiscal_year = to_fiscal_year(title)
    if fiscal_year is None:
        issues.append(f"{resource_name}: シートタイトルから年度を抽出できず: {title!r}")

    grid = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    rows_out = []
    current_year = None
    i = 0
    while i < len(grid):
        row = grid[i]
        b = row[1] if len(row) > 1 else None
        f = row[5] if len(row) > 5 else None  # "採水日" ラベル列
        if isinstance(b, str) and "月" in b and f == "採水日":
            month_label = b
            day_raw = row[6] if len(row) > 6 else None
            m = MONTH_HDR_RE.search(month_label)
            if not m:
                issues.append(f"{resource_name}: 月ラベルを解釈できず: {month_label!r}")
                i += 1
                continue
            if m.group(1):
                y = to_fiscal_year(m.group(1))
                if y is not None:
                    current_year = y
            month_num = int(m.group(2))
            # 次行が "測定項目" ヘッダのはず
            if i + 1 >= len(grid) or grid[i + 1][1] != "測定項目":
                issues.append(f"{resource_name} {month_label}: 直後に「測定項目」行が続かない")
                i += 1
                continue
            j = i + 2
            while j < len(grid) and grid[j][1] not in (None, ""):
                site_row = grid[j]
                site_raw = site_row[1]
                river_ja = re.sub(r"[　\s]+", "", str(site_raw)) if site_raw is not None else None
                # 採水日 -> 日付
                sampled_on = None
                sampled_on_raw = None
                note_bits = []
                if current_year is not None and day_raw:
                    day_str = str(day_raw)
                    dm = re.fullmatch(r"(\d+)日", day_str)
                    if dm:
                        try:
                            sampled_on = datetime.date(current_year, month_num, int(dm.group(1))).isoformat()
                        except ValueError:
                            note_bits.append(f"採水日の日付が不正: {current_year}-{month_num}-{dm.group(1)}")
                        sampled_on_raw = f"{current_year}年{month_num}月{day_str}"
                    else:
                        note_bits.append("採水日が複数日にまたがる(原文の範囲表記)ため単一日付を確定できない")
                        sampled_on_raw = f"{current_year}年{month_num}月{day_str}"
                for col_idx, variable_ja, unit_ja in ATSUGI_VARS:
                    if col_idx >= len(site_row):
                        continue
                    raw_val = site_row[col_idx]
                    value, value_raw, detection_flag = None, None, None
                    if raw_val is None or (isinstance(raw_val, str) and raw_val.strip() in ("", "－", "―", "-")):
                        value_raw = raw_val
                    elif isinstance(raw_val, (int, float)):
                        value = float(raw_val)
                        value_raw = str(raw_val)
                    else:
                        s = str(raw_val).strip()
                        dl = BELOW_DL_RE.match(s)
                        if dl:
                            value = None
                            value_raw = s
                            detection_flag = "<"
                        else:
                            num, kind = to_number(s)
                            value_raw = s
                            if kind in ("int", "float"):
                                value = float(num)
                            else:
                                note_bits.append(f"数値化できない値: {variable_ja}={s!r}")
                    rows_out.append({
                        "site_name_ja": river_ja,
                        "river_ja": river_ja,
                        "fiscal_year": fiscal_year,
                        "sampled_on": sampled_on,
                        "sampled_on_raw": sampled_on_raw,
                        "variable_ja": variable_ja,
                        "value": value,
                        "value_raw": value_raw,
                        "unit_ja": unit_ja,
                        "detection_flag": detection_flag,
                        "lat": None,
                        "lon": None,
                        "note_ja": ("地点名の詳細(住所等)は原表に記載なし。河川名を地点名として使用。"
                                     + ("" if not note_bits else " / ".join(note_bits))),
                        "source_id": "atsugi_river_water_quality",
                        "source_ref": f"{resource_url}#{month_label}:{site_raw}:{variable_ja}",
                    })
                j += 1
            i = j
        else:
            i += 1
    return rows_out, issues


def harvest_atsugi_river_water_quality(bodik_pkgs):
    target = None
    for p in bodik_pkgs:
        org = (p.get("organization") or {}).get("title")
        if org == "厚木市" and "河川水質調査結果" in (p.get("title") or ""):
            target = p
            break
    if target is None:
        register("atsugi_river_water_quality", "厚木市 河川水質調査結果(相模川ほか)", "厚木市",
                  BODIK_BASE, "河川水質", "CKAN resource download (XLSX)", "CSV",
                  "CC BY 4.0", True, 0,
                  "BODIK収穫結果内に厚木市「河川水質調査結果」パッケージが見つからなかった。取得できず。")
        print("  [skip] atsugi_river_water_quality: パッケージが見つからない")
        return

    resources = [r for r in (target.get("resources") or []) if (r.get("format") or "").upper() == "XLSX"]
    also_park = [p for p in bodik_pkgs
                 if (p.get("organization") or {}).get("title") == "厚木市" and p.get("title") == "都市公園一覧表"]

    all_rows = []
    all_issues = []
    fetched_years = []
    failed = []
    outdir = RAW / "atsugi_river_water_quality"
    for r in sorted(resources, key=lambda r: r.get("name") or ""):
        url = r.get("url")
        name = r.get("name") or url.rsplit("/", 1)[-1]
        fname = url.rsplit("/", 1)[-1]
        dest = outdir / fname
        try:
            download(url, dest)
        except Exception as e:
            failed.append((name, url, repr(e)))
            print(f"  [fail-download] {name}: {e}")
            continue
        try:
            rows, issues = parse_atsugi_xlsx(dest, url, name)
        except Exception as e:
            failed.append((name, url, f"parse error: {e!r}"))
            print(f"  [fail-parse] {name}: {e}")
            continue
        if issues:
            all_issues.extend(issues)
        if rows:
            all_rows.extend(rows)
            fys = sorted({r["fiscal_year"] for r in rows if r["fiscal_year"] is not None})
            fetched_years.extend(fys)
        else:
            failed.append((name, url, "行が0件(パース結果が空)"))

    write_jsonl("atsugi_river_water_quality", all_rows)
    df = pd.DataFrame(all_rows, columns=[
        "site_name_ja", "river_ja", "fiscal_year", "sampled_on", "sampled_on_raw",
        "variable_ja", "value", "value_raw", "unit_ja", "detection_flag",
        "lat", "lon", "note_ja", "source_id", "source_ref"])
    df.to_csv(PROC / "atsugi_river_water_quality.csv", index=False)

    notes = (f"厚木市「河川水質調査結果」(相模川・中津川・小鮎川・玉川、月毎シート)。"
             f"XLSX全{len(resources)}本(平成14〜令和2年度)のうち19本すべてが同一の表構造"
             f"(B1:H99, 月毎12ブロック×4河川×5項目)であることを1本ずつ確認したうえで統合した。"
             f"落とした年度: {'なし' if not failed else ', '.join(x[0] for x in failed)}。"
             f"定量下限表記(例: 「1未満」)は value=null / value_raw=原文 / detection_flag='<' として保持し0に潰していない。"
             f"採水日が範囲表記(例:「7～8日」、平成17年度以降の大半)の場合は単一日付を確定できないため sampled_on=null とし、"
             f"sampled_on_raw に原文相当を残した。緯度経度は原表に記載がないため null(推測していない)。"
             f"同一パッケージ内に「項目毎」ピボットシートがあるが「月毎」シートと同一データの別集計のため使用していない。"
             + (f" 同一団体の「都市公園一覧表」XLSX(1本)も確認したが本タスクの対象外のため未パース(register未登録)。"
                if also_park else ""))
    register("atsugi_river_water_quality", "厚木市 河川水質調査結果(相模川・中津川・小鮎川・玉川)", "厚木市",
              target.get("url") or f"{BODIK_BASE}/dataset/{target.get('name')}",
              "河川水質", "CKAN resource download (XLSX, openpyxl)", "CSV",
              "Creative Commons Attribution 4.0 International (CC BY 4.0)", True, len(all_rows), notes)

    print(f"  atsugi_river_water_quality: {len(all_rows)} rows  fiscal_years={sorted(set(fetched_years))}")
    if failed:
        print(f"  [failed resources] {failed}")
    if all_issues:
        print(f"  [parse issues] {len(all_issues)} 件: {all_issues[:10]}")
    return all_rows, failed, all_issues


# ============================================================
# 3. 横浜市 CKAN 目録収穫
# ============================================================

YOKOHAMA_BASE = "https://data.city.yokohama.lg.jp"


def harvest_yokohama():
    pkgs = ckan_harvest(YOKOHAMA_BASE)
    ds_rows, res_rows = ckan_to_rows(YOKOHAMA_BASE, "yokohama", pkgs)
    write_jsonl("ckan_datasets_yokohama", ds_rows)
    write_jsonl("ckan_resources_yokohama", res_rows)
    pd.DataFrame(ds_rows).to_csv(PROC / "ckan_datasets_yokohama.csv", index=False)
    pd.DataFrame(res_rows).to_csv(PROC / "ckan_resources_yokohama.csv", index=False)
    print(f"  yokohama: {len(pkgs)} datasets")
    register("ckan_yokohama", "横浜市オープンデータポータル", "横浜市", YOKOHAMA_BASE,
              "オープンデータ目録", "CKAN API (package_search 全件収穫)", "JSON",
              "各データセット個別(多くは CC BY 4.0)。https://data.city.yokohama.lg.jp/terms.html に"
              "「特別の記載がない限りクリエイティブ・コモンズ・ライセンス 表示 4.0 国際のもとでライセンス」との記載を確認。",
              True, len(pkgs),
              "package_search 全件収穫。generator: ckan 2.9.11 の正規CKAN。"
              "対象データの地理的注記: 本インスタンスに含まれる河川関連データ(鶴見川・帷子川・大岡川・侍従川・"
              "境川・宮川など)は相模川水系ではない(境川も相模湾に注ぐ別水系)。流域カルテの対象(相模川水系)外の"
              "隣接下流都市の参考データという位置づけであり、本体データと混同しないこと。")
    return pkgs, ds_rows, res_rows


# ============================================================
# 4. 横浜市 河川水位オープンデータ (yokohama_river_waterlevel)
# ============================================================

def harvest_yokohama_waterlevel(yoko_pkgs):
    target = None
    for p in yoko_pkgs:
        if "河川水位オープンデータ" in (p.get("title") or ""):
            target = p
            break
    if target is None:
        register("yokohama_river_waterlevel", "横浜市 河川水位オープンデータ", "横浜市(下水道河川局)",
                  YOKOHAMA_BASE, "河川水位", "CKAN resource download (CSV)", "CSV", "CC BY", True, 0,
                  "横浜市CKAN収穫結果内に「河川水位オープンデータ」パッケージが見つからなかった。取得できず。")
        print("  [skip] yokohama_river_waterlevel: パッケージが見つからない")
        return

    resources = target.get("resources") or []
    master_res = [r for r in resources if "マスター" in (r.get("name") or "")]
    monthly_res = [r for r in resources if r not in master_res]
    outdir = RAW / "yokohama_river_waterlevel"
    failed = []

    # -- 地点マスタ --
    id_map, lat_map, lon_map = {}, {}, {}
    if master_res:
        r = master_res[0]
        url = r.get("url")
        dest = outdir / url.rsplit("/", 1)[-1]
        try:
            download(url, dest)
            sdf = pd.read_csv(dest, encoding="utf-8-sig", dtype=str)
            sdf = sdf.rename(columns={
                "観測所番号": "station_id", "観測所": "station_name_ja", "水系": "water_system_ja",
                "河川名": "river_ja", "緯度": "lat", "経度": "lon"})
            sdf["source_id"] = "yokohama_river_waterlevel"
            sdf["source_ref"] = url
            sdf.to_csv(PROC / "yokohama_river_waterlevel_stations.csv", index=False)
            id_map = dict(zip(sdf["station_name_ja"], sdf["station_id"]))
            lat_map = {k: float(v) for k, v in zip(sdf["station_name_ja"], sdf["lat"]) if pd.notna(v)}
            lon_map = {k: float(v) for k, v in zip(sdf["station_name_ja"], sdf["lon"]) if pd.notna(v)}
            print(f"  yokohama_river_waterlevel_stations: {len(sdf)} rows")
        except Exception as e:
            failed.append(("水位地点マスターファイル", url, repr(e)))
            print(f"  [fail] 水位地点マスターファイル: {e}")
    else:
        failed.append(("水位地点マスターファイル", None, "リソースが見つからない"))

    # -- 月次データ --
    total_rows = 0
    n_files_ok = 0
    unmatched_stations = set()
    non_numeric_count = 0
    out_csv = PROC / "yokohama_river_waterlevel.csv"
    out_jsonl_path = PROC / "yokohama_river_waterlevel.jsonl"
    cols = ["station_id", "station_name_ja", "river_ja", "observed_at", "water_level_m",
            "value_raw", "lat", "lon", "source_id", "source_ref"]
    first_write = True
    jf = open(out_jsonl_path, "w", encoding="utf-8")
    for r in sorted(monthly_res, key=lambda r: r.get("name") or ""):
        url = r.get("url")
        name = r.get("name") or url.rsplit("/", 1)[-1]
        dest = outdir / url.rsplit("/", 1)[-1]
        try:
            download(url, dest)
        except Exception as e:
            failed.append((name, url, repr(e)))
            print(f"  [fail-download] {name}: {e}")
            continue
        try:
            mdf = pd.read_csv(dest, encoding="utf-8-sig", dtype=str)
        except Exception as e:
            failed.append((name, url, f"read_csv error: {e!r}"))
            print(f"  [fail-parse] {name}: {e}")
            continue
        mdf = mdf.rename(columns={"日付": "date_raw", "水位(cm)": "level_raw",
                                    "観測所名": "station_name_ja", "河川名": "river_ja"})
        # ベクトル化処理(行数が全24本合計で1000万行規模のため iterrows は使わない)
        this_unmatched = set(mdf["station_name_ja"].dropna().unique()) - set(id_map.keys())
        unmatched_stations |= this_unmatched
        odf = pd.DataFrame({
            "station_id": mdf["station_name_ja"].map(id_map),
            "station_name_ja": mdf["station_name_ja"],
            "river_ja": mdf["river_ja"],
            "observed_at": pd.to_datetime(mdf["date_raw"], format="%Y/%m/%d %H:%M", errors="coerce")
                             .dt.strftime("%Y-%m-%dT%H:%M:%S+09:00"),
            "value_raw": mdf["level_raw"],
            "lat": mdf["station_name_ja"].map(lat_map),
            "lon": mdf["station_name_ja"].map(lon_map),
        })
        wl = pd.to_numeric(mdf["level_raw"], errors="coerce")
        non_numeric_count += int((wl.isna() & mdf["level_raw"].notna()).sum())
        odf["water_level_m"] = (wl / 100.0).round(3)
        odf["source_id"] = "yokohama_river_waterlevel"
        odf["source_ref"] = url
        odf = odf[cols]
        odf.to_csv(out_csv, index=False, mode="w" if first_write else "a", header=first_write)
        # to_json(lines=True) は最終レコードの後に既に改行を1つ付けるため、余分な "\n" は足さない
        # (足すと月次ファイルの境目に空行が入り、JSONL の逐次パースを崩す)
        odf.to_json(jf, orient="records", lines=True, force_ascii=False)
        first_write = False
        total_rows += len(odf)
        n_files_ok += 1
        print(f"  [ok] {name}: {len(odf)} rows (累計 {total_rows})")
    jf.close()
    if total_rows == 0 and out_csv.exists():
        out_csv.unlink()

    notes = (f"水位地点マスター1本+月次データ{len(monthly_res)}本(令和6年度4月〜)。"
             f"5分間隔の実測水位(cm)を water_level_m(m換算)に変換して格納。"
             f"月次CSVの観測所名がマスター未収載だったもの: {sorted(unmatched_stations) if unmatched_stations else 'なし'}"
             f"(該当分は station_id/lat/lon が null。マスターファイルの更新が観測所追加に追いついていない可能性)。"
             f"数値化できなかった水位値: {non_numeric_count}件。"
             f"地理的注記: 対象河川は境川水系・帷子川水系・鶴見川水系等で相模川水系ではない(隣接下流都市の参考データ)。")
    register("yokohama_river_waterlevel", "横浜市 河川水位オープンデータ", "横浜市(下水道河川局)",
              target.get("url") or f"{YOKOHAMA_BASE}/dataset/{target.get('name')}",
              "河川水位", "CKAN resource download (CSV)", "CSV",
              "Creative Commons Attribution (CC BY)", True, total_rows, notes)
    print(f"  yokohama_river_waterlevel: {total_rows} rows across {n_files_ok}/{len(monthly_res)} monthly files")
    if failed:
        print(f"  [failed resources] {failed}")
    return total_rows, failed


# ============================================================
# 5. 横浜市 生物多様性(河川 生物指標水質評価) (yokohama_bio_indicator)
# ============================================================

PERIOD_RE = re.compile(r"^\d{4}-\d{4}年度$")


def parse_bio_indicator_csv(path, source_ref_base):
    df = pd.read_csv(path, encoding="utf-8-sig", header=1)
    period_cols = [c for c in df.columns if isinstance(c, str) and PERIOD_RE.match(c.strip())]
    quality_cols = [c for c in df.columns if isinstance(c, str) and "水質評価値" in c]
    quality_col = quality_cols[0] if quality_cols else None
    rows = []
    for _, row in df.iterrows():
        site_id = row.get("地点番号")
        river_ja = row.get("河川名")
        if pd.isna(river_ja):
            continue
        tributary = row.get("支川名")
        site_name_ja = row.get("地点名")
        water_area_class = row.get("水域区分")
        target_quality = row.get(quality_col) if quality_col else None
        note_bits = [f"地点番号: {site_id}", f"水域区分: {water_area_class}", f"水質評価値(達成目標): {target_quality}"]
        if isinstance(tributary, str) and tributary.strip() not in ("-", ""):
            note_bits.append(f"支川: {tributary}")
        for pcol in period_cols:
            result_raw = row.get(pcol)
            result_raw = None if pd.isna(result_raw) else str(result_raw).strip()
            achieved = None
            if result_raw == "○":
                achieved = 1
            elif result_raw == "×":
                achieved = 0
            fiscal_year = to_fiscal_year(pcol)
            rows.append({
                "river_ja": river_ja,
                "site_name_ja": site_name_ja,
                "fiscal_year": fiscal_year,
                "indicator_ja": "生物指標による水質評価",
                "result_raw": result_raw,
                "achieved": achieved,
                "note_ja": (f"評価期間: {pcol}(2ヵ年度分の集計。fiscal_yearは期間の開始年度) / " + " / ".join(note_bits)),
                "source_id": "yokohama_bio_indicator",
                "source_ref": f"{source_ref_base}:{site_id}:{pcol}",
            })
    return rows


def harvest_yokohama_bio_indicator(yoko_pkgs):
    target = None
    for p in yoko_pkgs:
        title = p.get("title") or ""
        if "環境管理計画年次報告書" in title and "資料編" in title:
            target = p
            break
    if target is None:
        register("yokohama_bio_indicator", "横浜市 環境管理計画年次報告書 資料編(生物多様性)", "横浜市(みどり環境局)",
                  YOKOHAMA_BASE, "生物多様性", "CKAN resource download (ZIP)", "CSV", "CC BY", True, 0,
                  "横浜市CKAN収穫結果内に「環境管理計画年次報告書 資料編」パッケージが見つからなかった。取得できず。")
        print("  [skip] yokohama_bio_indicator: パッケージが見つからない")
        return

    resources = target.get("resources") or []
    bio_res = [r for r in resources if "生物多様性" in (r.get("name") or "")]
    outdir = RAW / "yokohama_bio_indicator"
    failed = []
    all_rows = []
    if not bio_res:
        failed.append(("生物多様性ZIP", None, "リソースが見つからない"))
    else:
        r = bio_res[0]
        url = r.get("url")
        zdest = outdir / url.rsplit("/", 1)[-1]
        try:
            download(url, zdest)
        except Exception as e:
            failed.append(("生物多様性ZIP", url, repr(e)))
            print(f"  [fail-download] 生物多様性ZIP: {e}")
        else:
            extract_dir = outdir / "extracted"
            extract_dir.mkdir(parents=True, exist_ok=True)
            with zipfile.ZipFile(zdest) as zf:
                names = zf.namelist()
                zf.extractall(extract_dir)
            # 目的のCSV: 13-1_水域の生物調査結果（河川）抜粋_生物指標による水質評価の目標達成状況.csv
            csv_candidates = [n for n in names
                               if n.endswith(".csv") and "13-1" in n and "河川" in n and "目標達成状況" in n]
            if not csv_candidates:
                failed.append(("13-1河川生物指標CSV", url, f"ZIP内に対象CSVが見つからない。内訳: {names}"))
            else:
                inner = csv_candidates[0]
                csv_path = extract_dir / inner
                try:
                    all_rows = parse_bio_indicator_csv(csv_path, f"{url}#{inner}")
                except Exception as e:
                    failed.append(("13-1河川生物指標CSV", url, f"parse error: {e!r}"))
                    print(f"  [fail-parse] 13-1河川生物指標CSV: {e}")
            large_pdfs = [n for n in names if n.lower().endswith(".pdf")]
            print(f"  [note] ZIP内の大容量PDF{len(large_pdfs)}本は展開してrawに置くのみで表抽出はしていない: {large_pdfs}")

    write_jsonl("yokohama_bio_indicator", all_rows)
    df = pd.DataFrame(all_rows, columns=[
        "river_ja", "site_name_ja", "fiscal_year", "indicator_ja", "result_raw",
        "achieved", "note_ja", "source_id", "source_ref"])
    df.to_csv(PROC / "yokohama_bio_indicator.csv", index=False)

    rivers = sorted({r["river_ja"] for r in all_rows}) if all_rows else []
    notes = (f"「環境管理計画年次報告書 資料編」の「生物多様性」ZIP(約44MB)内、"
             f"13-1_水域の生物調査結果（河川）抜粋_生物指標による水質評価の目標達成状況.csv を展開・パース。"
             f"対象河川: {rivers}(鶴見川・帷子川・大岡川・侍従川・境川・宮川)。42地点×最大2評価期間"
             f"(2018-2019年度/2022-2023年度)を縦持ちに変換。achievedは○=1/×=0/該当なし(-)はnullで、"
             f"原文はresult_rawに保持。同ZIP内の生物調査結果PDF(河川22MB/海域18MB)は展開してrawに置くのみで"
             f"表抽出はしていない(別Tierの仕事)。"
             f"地理的注記: 対象河川は相模川水系ではない(境川も相模湾に注ぐ別水系)。相模川本流域への直接的な"
             f"追加価値は限定的で、隣接下流都市の参考データという位置づけ。")
    register("yokohama_bio_indicator", "横浜市 環境管理計画年次報告書 資料編(河川 生物指標水質評価)",
              "横浜市(みどり環境局)",
              target.get("url") or f"{YOKOHAMA_BASE}/dataset/{target.get('name')}",
              "生物多様性・水質", "CKAN resource download (ZIP→CSV)", "CSV",
              "Creative Commons Attribution (CC BY)", True, len(all_rows), notes)
    print(f"  yokohama_bio_indicator: {len(all_rows)} rows  rivers={rivers}")
    if failed:
        print(f"  [failed resources] {failed}")
    return all_rows, failed


# ============================================================
# main
# ============================================================

def main():
    print("== 1. BODIK(神奈川県参加団体) 目録収穫 ==")
    bodik_pkgs, _, _ = harvest_bodik()

    print("\n== 2. 厚木市 河川水質調査結果 (実データ) ==")
    harvest_atsugi_river_water_quality(bodik_pkgs)

    print("\n== 3. 横浜市 CKAN 目録収穫 ==")
    yoko_pkgs, _, _ = harvest_yokohama()

    print("\n== 4. 横浜市 河川水位オープンデータ (実データ) ==")
    harvest_yokohama_waterlevel(yoko_pkgs)

    print("\n== 5. 横浜市 生物多様性(河川 生物指標水質評価) (実データ) ==")
    harvest_yokohama_bio_indicator(yoko_pkgs)

    print("\n完了。")


if __name__ == "__main__":
    main()
