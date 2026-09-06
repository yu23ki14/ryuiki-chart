"""神奈川県 地盤沈下調査結果 (観測井戸/水準点 × 年度の縦持ち)
入口: https://www.pref.kanagawa.jp/docs/pf7/cnt/f41044/p81475.html

令和4〜6年度の3年分が掲載され、令和5・6年度は表(Excel)直リンク、3年度ともPDF直リンクがある。
令和3年度以前は「環境課水環境グループにお問い合わせください」と明記されており機械取得不可
(問合せ行為は収集規約で禁止のため対象外)。

Excel(r5table.xlsx / r6table.xlsx)は18シート構成。実際に開いて確認した結果:
- 表9「地下水位等の観測所の諸元」: 観測井戸の台帳(№, 観測所名, 所在地, 観測機関名(=市町名), 稼働状況等)。
- 表11-1/11-2-1/11-2-2/11-3「地下水位の経年変化」: 観測井戸×年度の地下水位(T.P.m)の縦持ちに最適な
  マトリクス表(行=年度, 列=観測井戸)。S55(1980)〜最新年度までの長期系列。
- 表4「沈下状況の経年変化（全域）」: 年度ごとの県全域集計(有効水準点数・沈下水準点数・調査面積・沈下面積)と
  その年の最大沈下地点(水準点№・所在地・沈下量)。S48(1973)〜最新年度。
- 表1/2/3/5/6/7/8/10-x/12 は「当該年度単年スナップショット」または「市町別集計で観測井戸粒度ではない」
  ため今回は対象外とした。特に表3「市町別最大沈下量（年間・累計）」は「累計最大沈下量」の列見出しが
  '調査開始以来'という同一ラベルで2列存在し(それぞれ異なる基準点・値を指す様子だが元シートの結合セル
  情報だけでは意味の切り分けを確認できなかった)、無理に構造化すると誤ったラベル付けになるリスクが高いため
  見送った(「無理に構造化しない」の方針に従う)。

r5table.xlsx と r6table.xlsx は両方とも「主データ」として指示されているため両方処理する。
実際に全シートをセル単位で突き合わせたところ、表4/表9/表11-xの経年系列は r6 が r5 の完全上位互換
(r6に最新年度の行が1行追加されているだけで、重複する年度の値は完全一致)であり、令和7年9月30日の
「令和4年と令和5年の一部データを修正しました」という更新履歴に対応する数値差異はこの3表では検出できなかった
(単年スナップショット表である表1/2/3/5/6/8/12は年度が異なり単純比較できないため確認していない)。
それでも両ファイルを重複除去せずにそのまま出力する(重複判定・優先順位付けはローダー側の仕事、というこの
プロジェクトの一貫した方針に従う)。source_ref でどちらのファイル由来かは常に判別できる。
"""
import sys, csv, re, pathlib
sys.path.insert(0, "scripts")
from common import get, download, register, write_jsonl, to_number, PROC, RAW

import openpyxl
from bs4 import BeautifulSoup

ENTRY_URL = "https://www.pref.kanagawa.jp/docs/pf7/cnt/f41044/p81475.html"
SID = "kanagawa_jiban_chinka"
RAW_DIR = RAW / SID
PUBLISHER = "神奈川県 環境農政局 環境部環境課 水環境グループ"

ERA_BASE = {"S": 1925, "Ｓ": 1925, "H": 1988, "Ｈ": 1988, "R": 2018, "Ｒ": 2018}
ERA_LABEL_RE = re.compile(r"^([SHRＳＨＲ])\s*(\d+|元)$")


def parse_year_label(label, state):
    """行ラベル(和暦、または元号なしの継続年数値)から西暦年度を求める。
    元号付きラベル('Ｓ55','H 1'等)が出るたびに state['base']を更新し、以降の
    素の数値ラベル(56,57,...)はその元号の続きとして解釈する(このExcelの一貫した書式)。
    """
    if isinstance(label, bool):
        return None
    if isinstance(label, (int, float)):
        if state.get("base") is None:
            return None
        return state["base"] + int(label)
    if isinstance(label, str):
        s = label.strip().replace("　", "")
        m = ERA_LABEL_RE.match(s)
        if m:
            base = ERA_BASE[m.group(1)]
            n = 1 if m.group(2) == "元" else int(m.group(2))
            state["base"] = base
            return base + n
        m2 = re.fullmatch(r"(19|20)\d{2}", s)
        if m2:
            return int(s)
    return None


def norm_well_id(v):
    s = str(v).strip()
    s = re.sub(r"(№|No\.|No,|NO\.)", "", s, flags=re.I)
    return s.strip()


def base_well_id(well_id):
    return re.sub(r"-\d+$", "", well_id)


def clean_cell(raw):
    """(value, value_raw) を返す。空欄/ダッシュ類は value=None。"""
    if raw is None:
        return None, None
    value_raw = str(raw)
    val, vtype = to_number(raw)
    if vtype == "string":
        val = None
    return val, value_raw


def find_header_row(ws, text, col=None, max_row=6):
    for r in range(1, min(max_row, ws.max_row) + 1):
        rng = range(1, ws.max_column + 1) if col is None else [col]
        for c in rng:
            v = ws.cell(r, c).value
            if isinstance(v, str) and text in v.replace("　", ""):
                return r, c
    return None, None


def sheet_by_prefix(wb, prefix):
    for sn in wb.sheetnames:
        if sn.strip().startswith(prefix):
            return wb[sn]
    return None


def collect_notes(ws):
    """シート内の「注N ...」テキストセルを全部集める(原文のまま)。"""
    notes = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and re.match(r"^注[0-9１-３]", c.value.strip()):
                notes.append(re.sub(r"\s+", " ", c.value.strip()))
    return notes


def note_for_row(sheet_notes, fiscal_year, well_name_ja):
    applicable = []
    for n in sheet_notes:
        if ("地震" in n) and fiscal_year == 2011:
            applicable.append(n)
        elif ("大原観測所" in n) and well_name_ja and "大原" in well_name_ja:
            applicable.append(n)
        elif ("最大沈下点" in n) and fiscal_year in (2019, 2023):
            applicable.append(n)
    return " / ".join(applicable) if applicable else None


MUNI_RE = re.compile(r"^(.+?(?:市|町|村))")


def extract_muni(address):
    if not address:
        return None
    first_line = str(address).split("\n")[0].split("　")[0]
    m = MUNI_RE.match(first_line)
    return m.group(1) if m else None


def parse_table9(ws, url, fname):
    """表9 観測所の諸元 -> 井戸台帳 dict と プロフィール行(rows)"""
    hdr_row, _ = find_header_row(ws, "観測所名")
    if hdr_row is None:
        print("  [warn] 表9: 「観測所名」ヘッダが見つからない")
        return {}, []
    data_start = hdr_row + 3  # №/種別subheader/単位subheaderの3行分
    lookup, rows = {}, []
    for r in range(data_start, ws.max_row + 1):
        no = ws.cell(r, 2).value
        if no is None:
            continue
        well_id = norm_well_id(no)
        name = ws.cell(r, 3).value
        addr = ws.cell(r, 4).value
        well_type = ws.cell(r, 5).value
        depth = ws.cell(r, 6).value
        caliber = ws.cell(r, 7).value
        strainer = ws.cell(r, 8).value
        datum = ws.cell(r, 9).value
        obs_start = ws.cell(r, 10).value
        agency = ws.cell(r, 11).value  # 観測機関名 = 市町名そのもの
        kind = ws.cell(r, 12).value
        status = ws.cell(r, 13).value
        muni = agency if agency else extract_muni(addr)
        name_clean = re.sub(r"\s+", "", str(name)) if name else None
        lookup[well_id] = {"name": name_clean, "municipality": muni}
        note = (f"所在地:{addr!r} 種別:{well_type} 深度:{depth}m 口径:{caliber}mm "
                f"ストレーナー位置:{strainer}m 水位基準面高:{datum}(T.P.m) "
                f"観測開始:{obs_start} 観測種類:{str(kind).replace(chr(10), '/') if kind else None} "
                f"稼働状況:{status}")
        rows.append({
            "well_id": well_id, "well_name_ja": name_clean, "municipality_ja": muni,
            "fiscal_year": None, "variable_ja": "観測所プロフィール",
            "value": None, "value_raw": status, "unit_ja": None,
            "lat": None, "lon": None, "note_ja": note,
            "source_id": SID, "source_ref": f"{url}#{fname}!表9!well{well_id}",
        })
    return lookup, rows


def parse_water_level_sheet(ws, sheet_key, well_lookup, url, fname):
    hdr_row, _ = find_header_row(ws, "区分", col=2)
    if hdr_row is None:
        print(f"  [warn] {sheet_key}: 「区分」ヘッダが見つからない")
        return []
    name_row = hdr_row + 1
    wells = []
    for c in range(3, ws.max_column + 1):
        wid_raw = ws.cell(hdr_row, c).value
        if wid_raw is None:
            continue
        well_id = norm_well_id(wid_raw)
        header_name = ws.cell(name_row, c).value
        header_name = re.sub(r"\s+", "", str(header_name)) if header_name else None
        wells.append((c, well_id, header_name))
    sheet_notes = collect_notes(ws)

    rows = []
    state = {"base": None}
    r = name_row + 1
    while r <= ws.max_row:
        label = ws.cell(r, 2).value
        if label is None:
            r += 1
            continue
        if isinstance(label, str) and re.match(r"^注[0-9１-３]", label.strip()):
            r += 1
            continue
        fy = parse_year_label(label, state)
        if fy is None:
            r += 1
            continue
        for c, well_id, header_name in wells:
            raw = ws.cell(r, c).value
            value, value_raw = clean_cell(raw)
            if value is None:
                continue
            lk = well_lookup.get(base_well_id(well_id), {})
            well_name_ja = lk.get("name") or header_name
            municipality_ja = lk.get("municipality")
            rows.append({
                "well_id": well_id, "well_name_ja": well_name_ja, "municipality_ja": municipality_ja,
                "fiscal_year": fy, "variable_ja": "地下水位(年平均)",
                "value": value, "value_raw": value_raw, "unit_ja": "T.P. m",
                "lat": None, "lon": None,
                "note_ja": note_for_row(sheet_notes, fy, well_name_ja),
                "source_id": SID, "source_ref": f"{url}#{fname}!{sheet_key}!col{c}_row{r}",
            })
        r += 1
    return rows


AGGREGATE_VARS = [
    (3, "有効水準点数", "点"),
    (4, "沈下水準点数(1cm以上2cm未満)", "点"),
    (5, "沈下水準点数(2cm以上)", "点"),
    (6, "調査面積", "km2"),
    (7, "沈下面積(1cm以上2cm未満)", "km2"),
    (8, "沈下面積(2cm以上)", "km2"),
]


def parse_subsidence_sheet(ws, sheet_key, url, fname):
    hdr_row, hdr_col = find_header_row(ws, "水準点№")
    if hdr_row is None:
        print(f"  [warn] {sheet_key}: 「水準点№」ヘッダが見つからない")
        return []
    col_no, col_addr, col_amount = hdr_col, hdr_col + 1, hdr_col + 2
    sheet_notes = collect_notes(ws)
    rows = []
    state = {"base": None}
    r = hdr_row + 1
    while r <= ws.max_row:
        label = ws.cell(r, 2).value
        if label is None:
            r += 1
            continue
        if isinstance(label, str) and re.match(r"^注[0-9１-３]", label.strip()):
            r += 1
            continue
        fy = parse_year_label(label, state)
        if fy is None:
            r += 1
            continue
        ref_base = f"{url}#{fname}!{sheet_key}!row{r}"
        for c, var_ja, unit in AGGREGATE_VARS:
            raw = ws.cell(r, c).value
            value, value_raw = clean_cell(raw)
            if value is None:
                continue
            rows.append({
                "well_id": None, "well_name_ja": None, "municipality_ja": None,
                "fiscal_year": fy, "variable_ja": var_ja,
                "value": value, "value_raw": value_raw, "unit_ja": unit,
                "lat": None, "lon": None,
                "note_ja": note_for_row(sheet_notes, fy, None),
                "source_id": SID, "source_ref": ref_base + f"_col{c}",
            })
        point_no = ws.cell(r, col_no).value
        addr = ws.cell(r, col_addr).value
        amount_raw = ws.cell(r, col_amount).value
        amount, amount_raw_str = clean_cell(amount_raw)
        if point_no is not None and amount is not None:
            rows.append({
                "well_id": norm_well_id(point_no), "well_name_ja": None, "municipality_ja": None,
                "fiscal_year": fy, "variable_ja": "最大沈下量(基準点)",
                "value": amount, "value_raw": amount_raw_str, "unit_ja": "cm",
                "lat": None, "lon": None,
                "note_ja": " / ".join(x for x in [
                    f"水準点所在地(原文): {addr}" if addr else None,
                    note_for_row(sheet_notes, fy, None),
                ] if x),
                "source_id": SID, "source_ref": ref_base + f"_col{col_amount}",
            })
        r += 1
    return rows


WATER_LEVEL_SHEETS = ["表11-1", "表11-2-1", "表11-2-2", "表11-3"]


def process_workbook(xlsx_path, url, fname):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    rows = []
    ws9 = sheet_by_prefix(wb, "表9")
    well_lookup, profile_rows = ({}, [])
    if ws9 is not None:
        well_lookup, profile_rows = parse_table9(ws9, url, fname)
        rows += profile_rows
        print(f"  [{fname}] 表9 観測井戸台帳: {len(well_lookup)} 件, プロフィール行 {len(profile_rows)}")
    else:
        print(f"  [warn] {fname}: 表9が見つからない")

    for key in WATER_LEVEL_SHEETS:
        ws = sheet_by_prefix(wb, key)
        if ws is None:
            print(f"  [warn] {fname}: {key} が見つからない")
            continue
        wl_rows = parse_water_level_sheet(ws, key, well_lookup, url, fname)
        rows += wl_rows
        print(f"  [{fname}] {key} 地下水位: {len(wl_rows)} 行")

    ws4 = sheet_by_prefix(wb, "表4")
    if ws4 is not None:
        sub_rows = parse_subsidence_sheet(ws4, "表4", url, fname)
        rows += sub_rows
        print(f"  [{fname}] 表4 沈下状況(全域): {len(sub_rows)} 行")
    else:
        print(f"  [warn] {fname}: 表4が見つからない")

    return rows


def fetch_links():
    r = get(ENTRY_URL)
    r.encoding = r.apparent_encoding
    soup = BeautifulSoup(r.text, "html.parser")
    links = {}
    for a in soup.find_all("a", href=True):
        href = a["href"]
        m = re.search(r"/([a-zA-Z0-9_]+\.(?:xlsx|pdf))$", href)
        if m:
            fname = m.group(1)
            full = href if href.startswith("http") else "https://www.pref.kanagawa.jp" + href
            links[fname] = full
    return links


def main():
    links = fetch_links()
    print(f"[fetch] {ENTRY_URL}")
    print(f"  検出リンク: {links}")
    expected = ["r6_1.pdf", "r6table.xlsx", "r5.pdf", "r5table.xlsx", "r4.pdf"]
    missing = [e for e in expected if e not in links]
    if missing:
        print(f"  [warn] 想定リンクが見つからない: {missing} (ページ構成が変わった可能性)")

    downloaded = {}
    for fname, url in links.items():
        dest = RAW_DIR / fname
        download(url, dest)
        downloaded[fname] = (dest, url)
        print(f"  [download] {fname} <- {url} ({dest.stat().st_size} bytes)")

    all_rows = []
    for fname in ("r5table.xlsx", "r6table.xlsx"):
        if fname not in downloaded:
            print(f"  [skip] {fname} が取得できていない")
            continue
        dest, url = downloaded[fname]
        rows = process_workbook(dest, url, fname)
        all_rows += rows
        print(f"  [{fname}] 合計 {len(rows)} 行")

    fields = ["well_id", "well_name_ja", "municipality_ja", "fiscal_year", "variable_ja",
              "value", "value_raw", "unit_ja", "lat", "lon", "note_ja", "source_id", "source_ref"]
    p = PROC / f"{SID}.csv"
    with open(p, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for row in all_rows:
            w.writerow(row)
    print(f"  [write] {p}  {len(all_rows)} rows")
    write_jsonl(SID, all_rows)

    pdf_notes = []
    for fname in ("r4.pdf", "r5.pdf", "r6_1.pdf"):
        if fname in downloaded:
            dest, url = downloaded[fname]
            pdf_notes.append(f"{fname}({dest.stat().st_size}bytes, {url})")

    register(
        source_id=SID,
        name="地盤沈下調査結果(令和4〜6年度、観測井戸・水準点)",
        publisher=PUBLISHER,
        url=ENTRY_URL,
        category="地盤環境・地下水",
        access_method="入口HTMLから動的にExcel/PDF直リンクを取得。Excel(表9/表11-1,2-1,2-2,3/表4)をopenpyxlでパース",
        fmt="xlsx->csv+jsonl (PDFは raw 保存のみ)",
        license_="神奈川県サイトポリシー(出典明示で利用可、要確認)。CKAN上の当該データセットはCC-BY表記。",
        redistributable=1,
        record_count=len(all_rows),
        notes=(
            "令和3年度以前は「環境課水環境グループにお問い合わせください」と明記されており問合せ行為が"
            "必要なため取得していない(収集規約で問合せ操作は禁止)。"
            "入口ページに更新履歴として原文『令和7年9月30日に令和4年と令和5年の地盤沈下調査結果の一部"
            "データを修正しました。』とあるが、実際にr5table.xlsxとr6table.xlsxの表4/表9/表11-x(経年系列)を"
            "全セル突き合わせた限りでは、重複する年度の数値に差異は検出できなかった(修正は主に単年スナップ"
            "ショット表(表1/2/3/5/6/8/12、今回は対象外)側にあった可能性がある。今回対象の経年系列表では"
            "確認できていない、という限定的な結果である点に注意)。"
            "表11-1/11-3に『注2 平成23年の調査結果には、平成23年3月11日の東北地方太平洋沖地震が影響している"
            "ものと考えられる。』、表11-3に『注3 大原観測所の標高値算出のための水準測量値を平成29年結果より"
            "変更したため、平成28年結果との継続性はない。』という時系列比較を妨げる注記があり、該当行の"
            "note_ja に転記済み。表4にも同旨の『注1 平成23年の沈下状況には、平成23年3月11日の東北地方太平洋"
            "沖地震が影響しているものと考えられる。』と『注2 令和元年、令和５年の最大沈下点は、隔年調査地域の"
            "沈下量を1/2として比較した結果を表示。』があり同様に転記済み。"
            "表3(市町別最大沈下量、年間・累計)は「累計最大沈下量」列見出しが『調査開始以来』のまま2列存在し"
            "元シートの結合セル情報だけでは2列の意味の切り分けを確認できなかったため、無理に構造化せず"
            "対象外とした(raw のExcelには残っている)。"
            "表9(観測所諸元)は観測井戸の静的プロフィールとして fiscal_year=None・"
            "variable_ja='観測所プロフィール' で1井戸1行出力し、所在地・井戸諸元・稼働状況等はnote_jaに"
            "原文のまま格納した。緯度経度は原表に記載がないため常にnull(ジオコーディングによる補完はしていない)。"
            "PDF(r4.pdf/r5.pdf/r6_1.pdf)は data/raw/kanagawa_jiban_chinka/ に保存済みだが表抽出はしていない: "
            + "; ".join(pdf_notes)
        ),
    )


if __name__ == "__main__":
    main()
