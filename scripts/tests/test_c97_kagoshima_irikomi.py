"""c97（奄美群島の入込客・入域客）の解析・検算・cells の形のテスト。合成の小さな XLSX だけを使う（ネットワーク・DB なし）。"""
import openpyxl
import pytest

import c97_kagoshima_irikomi as c


def make_xlsx(path, years, bump=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tid, (sheet, _) in c.SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.append([None, "海路", "空路"])
        for k, y in enumerate(sorted(years, reverse=True)):
            base = 1000 * (y - 2000) + (500 if tid == "x1" else 0)   # 入込 > 入域
            ws.append([y, base + 1, base + 2 + (bump or 0 if (tid, y) == ("x1", 2022) else 0)])
    wb.save(path)
    return path


def test_parse_and_build(tmp_path):
    main = c.parse_workbook(make_xlsx(tmp_path / "m.xlsx", range(2005, 2025)))
    other = c.parse_workbook(make_xlsx(tmp_path / "o.xlsx", range(2005, 2023)))
    assert c.check(main, other) == 2 * 2 * 18
    rows = c.build_rows(main, other, "2026-10-10T00:00:00")
    assert len(rows) == 2 * 20 * 2
    assert len({(r["table_id"], r["row_key"]) for r in rows}) == 4
    r = next(r for r in rows if (r["table_id"], r["row_key"], r["fiscal_year"]) == ("x1", "奄美群島|空路", 2024))
    assert (r["value"], r["unit"], r["source_bbox"], r["col_key"], r["era_raw"]) == ("24502", "人", "入込客（人）!C2", "2024", "2024")
    assert r["verified_by"] == "auto:xlsx"
    assert next(r for r in rows if r["fiscal_year"] == 2022)["verified_by"] == "auto:xversion"
    assert not any(r["is_total"] for r in rows)       # 海路＋空路の合計は作らない
    # 入込客が入域客より大きい（notes の記述の前提）
    by = {(r["table_id"], r["row_key"], r["fiscal_year"]): int(r["value"]) for r in rows}
    assert all(by[("x1", k, y)] > by[("x2", k, y)] for k in ("奄美群島|海路", "奄美群島|空路") for y in c.YEARS)


def test_versions_that_disagree_stop(tmp_path):
    main = c.parse_workbook(make_xlsx(tmp_path / "m.xlsx", range(2005, 2025)))
    other = c.parse_workbook(make_xlsx(tmp_path / "o.xlsx", range(2005, 2023), bump=1))
    with pytest.raises(ValueError, match="食い違う"):
        c.check(main, other)


def test_missing_year_stops(tmp_path):
    main = c.parse_workbook(make_xlsx(tmp_path / "m.xlsx", [y for y in range(2005, 2025) if y != 2010]))
    with pytest.raises(ValueError, match="連続"):
        c.check(main, {"x1": {}, "x2": {}})


def test_non_integer_value_stops(tmp_path):
    p = make_xlsx(tmp_path / "m.xlsx", range(2005, 2025))
    wb = openpyxl.load_workbook(p)
    wb["入込客（人）"]["B3"] = "1,234"
    wb.save(p)
    with pytest.raises(ValueError, match="B3"):
        c.parse_workbook(p)


def test_notes_say_not_amami_island():
    kinds = {n["kind"] for n in c.NOTES}
    assert kinds == {"survey_scope", "footnote", "definition_change"}
    assert "奄美大島だけの値ではない" in next(n for n in c.NOTES if n["kind"] == "survey_scope")["text"]
