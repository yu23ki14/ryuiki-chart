"""奄美群島の入込客数・入域客数（BODIK 460001_guntou_irikomi_nyuiki, CC BY 4.0）を行政文書の cells に入れる。

設計: docs/plans/AMAMI_STEP2C.md §3c・§0 決定1・2。

- 取得: BODIK の package_show で XLSX のリソースを引き、2024年版（暦年 2005〜2024）を入れる。2022年版（2005〜2022）は、
  重なる18年の値が2024年版と一致することを確かめるだけに使う（食い違えば止まる）。呼び出しの間は 5 秒以上空ける。
- 読み方: openpyxl。シートは「入込客（人）」（table_id `x1`）と「入域客（人）」（`x2`）、A列=年、B列=海路、C列=空路。
- cells: row_key は `奄美群島|海路`・`奄美群島|空路`。col_key と era_raw は年（西暦の4桁）、fiscal_year には**暦年**を入れる
  （系列の年軸）。unit は「人」。source_bbox は `入込客（人）!B2` の形。海路と空路の合計は原本に無いので作らない。
- **奄美大島の値ではない**。奄美群島（5島・12市町村）全体の値。notes に書く（決定2）。
- 書き込みは doccells.py（自分の doc_id の行だけを入れ替える）。検算（年の連続・値の型・2版の一致）は書く前に済ませる。

使い方: python3 scripts/c97_kagoshima_irikomi.py [--no-fetch] [--no-register]
"""
import argparse
import pathlib
import re
import sys
import time

sys.path.insert(0, str(pathlib.Path(__file__).parent))
import common
import doccells

PKG = "460001_guntou_irikomi_nyuiki"
BASE = "https://data.bodik.jp"
SLEEP = 5.0   # BODIK の呼び出しの間（秒。続けて呼ぶと 403 になる）
RAW = common.RAW / "bodik_irikomi"
SOURCE_ID = "kagoshima_irikomi_amami"
DOC_ID = "bodik_460001_guntou_irikomi_2024"
MAIN_FILE, CHECK_FILE = "2024_guntou_irikomi_nyuiki.xlsx", "2022_guntou_irikomi_nyuiki.xlsx"
YEARS = range(2005, 2025)
ROUTES = ("海路", "空路")
SHEETS = {"x1": ("入込客（人）", 1), "x2": ("入域客（人）", 2)}   # table_id -> (シート名, page_no)
LICENSE = "CC BY 4.0"
EXTRACTOR = "openpyxl"
SCRIPT_ID = "c97"


def fetch(force=False):
    """package_show -> RAW に2つの XLSX を保存。-> {ファイル名: URL}"""
    urls = {}
    time.sleep(SLEEP)
    pkg = common.get_json(f"{BASE}/api/3/action/package_show", params={"id": PKG})["result"]
    for r in pkg["resources"]:
        name = r["url"].rsplit("/", 1)[-1]
        if name in (MAIN_FILE, CHECK_FILE):
            urls[name] = r["url"]
    if set(urls) != {MAIN_FILE, CHECK_FILE}:
        raise SystemExit(f"package_show に期待のリソースが無い: {sorted(urls)}")
    for name, url in urls.items():
        if (RAW / name).exists() and not force:
            continue
        time.sleep(SLEEP)
        common.download(url, RAW / name)
    return urls


def parse_workbook(path):
    """XLSX -> {table_id: {年: {経路: (値, セル番地)}}}。値は整数。読めないセルがあれば止める。"""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for tid, (sheet, _page) in SHEETS.items():
        if sheet not in wb.sheetnames:
            raise ValueError(f"{path}: シート {sheet!r} が無い（{wb.sheetnames}）")
        ws = wb[sheet]
        head = [ws.cell(1, c).value for c in (2, 3)]
        if tuple(head) != ROUTES:
            raise ValueError(f"{path} {sheet}: 見出しが {head}（期待 {list(ROUTES)}）")
        t = {}
        for r in range(2, ws.max_row + 1):
            y = ws.cell(r, 1).value
            if y is None:
                continue
            if not isinstance(y, int):
                raise ValueError(f"{path} {sheet}!A{r}: 年が整数でない: {y!r}")
            row = {}
            for c, route in zip((2, 3), ROUTES):
                v = ws.cell(r, c).value
                if not isinstance(v, int) or isinstance(v, bool) or v < 0:
                    raise ValueError(f"{path} {sheet}!{ws.cell(r, c).coordinate}: 値が非負の整数でない: {v!r}")
                row[route] = (v, ws.cell(r, c).coordinate)
            if y in t:
                raise ValueError(f"{path} {sheet}: 年 {y} が重複")
            t[y] = row
        out[tid] = t
    return out


def check(main, other):
    """検算。(1) 2024年版は暦年 2005〜2024 を欠かさず持つ。(2) 2022年版との重なりの値が一致する。"""
    for tid, t in main.items():
        if sorted(t) != list(YEARS):
            raise ValueError(f"{tid}: 年が {sorted(t)}（期待 2005〜2024 の連続）")
    n = 0
    for tid, t in other.items():
        for y, row in t.items():
            for route, (v, _) in row.items():
                if main[tid][y][route][0] != v:
                    raise ValueError(f"{tid} {y} {route}: 2024年版 {main[tid][y][route][0]} と 2022年版 {v} が食い違う")
                n += 1
    return n


def build_rows(main, other, extracted_at):
    rows = []
    for tid, (sheet, page) in SHEETS.items():
        for y in sorted(main[tid], reverse=True):
            for route in ROUTES:
                v, cell = main[tid][y][route]
                both = y in other[tid]
                rows.append(dict(
                    page_no=page, table_id=tid, row_key=f"奄美群島|{route}", col_key=str(y),
                    value_raw=str(v), value=str(v), value_type="int", unit="人", fiscal_year=y, era_raw=str(y),
                    source_text=f"{y}年 {route} {v:,}", source_bbox=f"{sheet}!{cell}", notes_ref=None,
                    is_total=0, merged=0, unreadable_reason=None, confidence=1.0,
                    extractor=EXTRACTOR, verified_by="auto:xversion" if both else "auto:xlsx", extracted_at=extracted_at))
    return rows


NOTES = [
    dict(kind="survey_scope", table_ids=["x1", "x2"], text=(
        "奄美群島（奄美大島・喜界島・徳之島・沖永良部島・与論島）全体の値であり、奄美大島だけの値ではない。"
        "島別・市町村別の内訳は原本に無い。")),
    dict(kind="footnote", table_ids=["x1", "x2"], text=(
        "年は暦年（1〜12月）で、年度ではない。BODIK のリソースの説明には「平成17年度から令和4年度まで」とあるが、"
        "表の年は西暦の4桁（2005〜）で、fiscal_year 列には暦年をそのまま入れている。")),
    dict(kind="definition_change", table_ids=["x1", "x2"], text=(
        "入込客数と入域客数は別のシートの数値で、どの年も入込客数が入域客数より大きい（例: 2024年の空路は入込 663,856 人・入域 553,097 人）。"
        "二つの定義（数え方の違い）は、このデータセットの説明にも原本の表にも書かれていないので、同じ量として足し引きしない。"
        "海路と空路の合計は原本に無く、cells にも作っていない。")),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-fetch", action="store_true", help="取得せず、RAW にある XLSX を使う")
    ap.add_argument("--no-register", action="store_true", help="source_registry に登録しない")
    a = ap.parse_args()
    urls = {} if a.no_fetch else fetch()
    main_p, other_p = RAW / MAIN_FILE, RAW / CHECK_FILE
    main_t, other_t = parse_workbook(main_p), parse_workbook(other_p)
    n = check(main_t, other_t)
    print(f"  [check] 2022年版との重なり {n} 値が一致")
    extracted_at = common.now()
    rows = build_rows(main_t, other_t, extracted_at)
    page = f"{BASE}/dataset/{PKG}"
    sha = common.sha256(main_p)
    con = common.cellsdb()
    try:
        doccells.write_doc(
            con, DOC_ID,
            document=dict(title="奄美群島入込客・入域客数（海路・空路、2005〜2024年）",
                          publisher="鹿児島県 大島支庁総務企画課（BODIK）", url=urls.get(MAIN_FILE) or page,
                          local_path=str(main_p.relative_to(common.ROOT)), doc_sha256=sha, n_pages=len(SHEETS),
                          fiscal_year=2024, license=LICENSE),
            cells=[dict(r, doc_sha256=sha) for r in rows], notes=NOTES,
            log=[dict(verdict="pass", note=f"{SCRIPT_ID}: 年の連続 2005〜2024・値は非負の整数・2022年版との一致 {n} 値")])
    finally:
        con.close()
    print(f"  [cells] {DOC_ID}: {len(rows)} 行、系列 {len({(r['table_id'], r['row_key']) for r in rows})}")
    if not a.no_register:
        common.register(SOURCE_ID, "奄美群島入込客・入域客数（海路・空路）", "鹿児島県 大島支庁総務企画課（BODIK）", page,
                        "観光", "HTTP GET (CKAN package_show -> XLSX)", "XLSX", LICENSE, True, len(rows),
                        "奄美群島全体の暦年の値（奄美大島だけではない）。入込客・入域客×海路・空路。")


if __name__ == "__main__":
    main()
