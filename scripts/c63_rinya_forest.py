"""林野庁 森林資源 -> data/processed/{rinya_forest_stats_prefecture, rinya_lidar_kanagawa_status}.{csv,jsonl}

## 1. 都道府県別 森林率・人工林率
出典一覧ページ: https://www.rinya.maff.go.jp/j/keikaku/genkyou/index2.html
  -> 最新版へのリンク: https://www.rinya.maff.go.jp/j/keikaku/genkyou/r4/1.html
     (見出し: 「都道府県別森林率・人工林率(令和4年3月31日現在)」)
  -> 全データダウンロード Excel: https://www.rinya.maff.go.jp/j/keikaku/genkyou/r4/attach/xls/1-1.xlsx
     シート「森林率・人工林率」(B1:K57) を使用。他の2シート(森林率Ｂ集計/人工林率Ｂ集計)は
     別の集計方法によるより詳細な表であり、本タスクの対象(都道府県別 森林率・人工林率の基本表)ではないため未使用。

原表記の注記(原文引用、要約しない):
  ※1 国土面積は、全国市町村要覧令和４年版による。
  ※2 全国及び北海道の森林率は北方領土を除いて算出した。

数値の扱い:
  Excelのセルは面積が表示形式「#,##0」(カンマ区切り整数表示)、森林率/人工林率が表示形式「0%」
  (整数%表示)だが、これはあくまで「表示」上の丸めであり、セルの実値(data_only=Trueで取得できる計算結果)
  はより高精度(例: 0.7059485118020843)。本スクリプトはセルの実値をそのまま*_raw列(文字列化)に残し、
  面積はha単位の数値、比率は0-1のratioをそのまま数値化した上で*_pctとして100倍した値も併記する
  (どちらも実値からの単純な単位変換であり、推測による値の補完ではない)。

## 2. 航空レーザ計測データの公開状況(神奈川県)
出典: https://www.rinya.maff.go.jp/j/keikaku/smartforest/smart_forestry.html
  (このURLはUTF-8で応答するがContent-Typeヘッダにcharset指定が無く、requestsの自動判定が
   ISO-8859-1に化けるため、r.content.decode("utf-8")で明示的にデコードする必要があった)

ページ構成:
  - 「申請に基づく提供データ」(#teikyou): 林野庁(本庁部局)が取得した航空レーザ測量成果等について、
    利用目的等に応じて個別に提供。提供対象は国の機関・地方公共団体・独立行政法人等の研究機関等に限定
    (原文: 「(ア)国の機関及び地方公共団体が行う業務」「(イ)独立行政法人、地方独立行政法人、国立大学法人、
    学校法人等の研究機関、公益社団法人及び公益財団法人が行う災害対策...」等)。問合せ先は
    「1.石川県能登地域」「2.その他」の2区分のみで、神奈川県固有の記載なし。
  - 「公開データ」(#koukai): 原文「林野庁保有の公開データについては、G空間情報センターの林野庁ページ
    (外部リンク)を参照願います。また、林野庁事業等により、都道府県によるオープンデータ化も進めており、
    併せて、G空間情報センターの都道府県ページもご覧ください。」
    -> つまり林野庁サイト自体には都道府県別(神奈川県含む)の公開データ一覧・地図は存在せず、
       すべてG空間情報センター(https://www.geospatial.jp/ckan/organization/rinya 等)経由。
       G空間情報センターは本タスクでは別の並行エージェントが担当するデータソースのため、
       契約上の重複回避のためこちらからは取得・深追いしない。
  - 令和5年度委託事業でのWEB-GIS実証(栃木県・兵庫県・高知県の航空レーザ解析データ)の記載はあるが、
    神奈川県は対象に含まれていない(原文: 「令和5年度委託事業では、栃木県、兵庫県及び高知県の
    航空レーザ解析データについて、WEB-GISでの配信を実証しました」)。

結論: 林野庁サイト上では神奈川県の航空レーザ計測データの公開状況を直接判断できる記述がない。
      「公開状況不明(林野庁サイト上は都道府県別の一覧なし。公開データの実体はG空間情報センター
      経由とのみ記載され、同センターは別エージェント担当のため本タスクでは未確認)」として記録する。
"""
import sys, re
sys.path.insert(0, "scripts")
from common import get, download, register, write_jsonl, to_fiscal_year, to_number, PROC, RAW

import openpyxl
import csv


def fetch_prefecture_forest_stats():
    source_id = "rinya_forest_stats_prefecture"

    # 1. 一覧ページから最新年度のリンクを抽出
    idx_url = "https://www.rinya.maff.go.jp/j/keikaku/genkyou/index2.html"
    r = get(idx_url)
    html = r.content.decode("utf-8", errors="replace") if r.encoding and r.encoding.lower() == "iso-8859-1" else r.text
    m = re.search(r'<h3><a href="(https://www\.rinya\.maff\.go\.jp/j/keikaku/genkyou/[^"]+/1\.html)">([^<]+)</a></h3>', html)
    if not m:
        register(source_id, "都道府県別森林率・人工林率", "林野庁", idx_url,
                  "forest", "html_scrape+excel", "xlsx", "政府標準利用規約2.0(想定)", 1, 0,
                  notes="一覧ページから最新年度リンクを抽出できず断念。")
        return
    detail_url, detail_label = m.group(1), m.group(2)
    print("latest year page:", detail_url, detail_label)

    r2 = get(detail_url)
    html2 = r2.content.decode("utf-8", errors="replace") if r2.encoding and r2.encoding.lower() == "iso-8859-1" else r2.text
    m2 = re.search(r'<a href="(\./attach/xls/[^"]+\.xlsx?)">([^<]*全データ[^<]*)</a>', html2)
    if not m2:
        register(source_id, "都道府県別森林率・人工林率", "林野庁", detail_url,
                  "forest", "html_scrape+excel", "xlsx", "政府標準利用規約2.0(想定)", 1, 0,
                  notes=f"{detail_url} からExcelへのリンクを抽出できず断念。")
        return
    xlsx_rel = m2.group(1)
    xlsx_url = re.sub(r'/[^/]+$', '/', detail_url) + xlsx_rel.lstrip("./")
    print("xlsx url:", xlsx_url)

    # h1見出しから年度表記を取得(例: 「令和4年3月31日現在」)
    h1m = re.search(r'<h1>([^<]+)</h1>', html2)
    title_ja = h1m.group(1) if h1m else detail_label
    fy_raw_m = re.search(r'[（(]([^)）]+)[)）]', title_ja)
    fy_raw = fy_raw_m.group(1) if fy_raw_m else title_ja
    fiscal_year = to_fiscal_year(fy_raw)

    dest = RAW / "rinya_forest_stats" / xlsx_url.split("/")[-1]
    dest_year_dir = RAW / "rinya_forest_stats"
    dest = dest_year_dir / f"{detail_url.rstrip('/').split('/')[-2]}_{xlsx_url.split('/')[-1]}"
    download(xlsx_url, dest)

    wb = openpyxl.load_workbook(dest, data_only=True)
    sheet_name = "森林率・人工林率"
    if sheet_name not in wb.sheetnames:
        register(source_id, "都道府県別森林率・人工林率", "林野庁", xlsx_url,
                  "forest", "excel", "xlsx", "政府標準利用規約2.0(想定)", 1, 0,
                  notes=f"シート'{sheet_name}'が見つからず断念。シート一覧: {wb.sheetnames}")
        return
    ws = wb[sheet_name]

    rows_out = []
    footnotes = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
        vals = [c.value for c in row]
        b, c_, d, e, f, h, j = vals[1], vals[2], vals[3], vals[4], vals[5], vals[7], vals[9]
        # 注記行 (B列に「※」始まりの文字列)
        if isinstance(b, str) and b.strip().startswith("※"):
            footnotes.append(b.strip())
            continue
        # 都道府県データ行: 番号(int)+都道府県名 または 「全  国」合計行
        pref_code = None
        pref_name = None
        if isinstance(b, int) and isinstance(c_, str):
            pref_code, pref_name = b, c_.strip()
        elif isinstance(b, str) and "全" in b and "国" in b:
            pref_name = re.sub(r"\s+", "", b.strip())
        else:
            continue
        if d is None and e is None and f is None:
            continue

        forest_area_raw = None if d is None else str(d)
        planted_area_raw = None if e is None else str(e)
        land_area_raw = None if f is None else str(f)
        forest_rate_raw = None if h is None else str(h)
        planted_rate_raw = None if j is None else str(j)

        forest_area_ha, _ = to_number(forest_area_raw)
        planted_forest_area_ha, _ = to_number(planted_area_raw)
        national_land_area_ha, _ = to_number(land_area_raw)
        forest_rate_ratio, _ = to_number(forest_rate_raw)
        planted_forest_rate_ratio, _ = to_number(planted_rate_raw)

        forest_rate_pct = round(forest_rate_ratio * 100, 4) if isinstance(forest_rate_ratio, (int, float)) else None
        planted_forest_rate_pct = round(planted_forest_rate_ratio * 100, 4) if isinstance(planted_forest_rate_ratio, (int, float)) else None

        notes_ja = ""
        if pref_name in ("北海道",) or (pref_code is None and pref_name and "全" in pref_name):
            notes_ja = "※2 全国及び北海道の森林率は北方領土を除いて算出した。"

        rows_out.append({
            "source_id": source_id,
            "source_ref": xlsx_url,
            "prefecture_ja": pref_name,
            "prefecture_code": pref_code,
            "forest_area_raw": forest_area_raw,
            "forest_area_ha": forest_area_ha,
            "planted_forest_area_raw": planted_area_raw,
            "planted_forest_area_ha": planted_forest_area_ha,
            "national_land_area_raw": land_area_raw,
            "national_land_area_ha": national_land_area_ha,
            "forest_rate_raw": forest_rate_raw,
            "forest_rate_ratio": forest_rate_ratio,
            "forest_rate_pct": forest_rate_pct,
            "planted_forest_rate_raw": planted_rate_raw,
            "planted_forest_rate_ratio": planted_forest_rate_ratio,
            "planted_forest_rate_pct": planted_forest_rate_pct,
            "fiscal_year_raw": fy_raw,
            "fiscal_year": fiscal_year,
            "notes_ja": notes_ja,
        })

    all_notes = " / ".join(footnotes)
    print(f"footnotes: {all_notes}")

    fieldnames = list(rows_out[0].keys())
    csv_path = PROC / f"{source_id}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows_out)
    print(f"  [write] {csv_path} {len(rows_out)} rows")
    write_jsonl(source_id, rows_out)

    register(source_id, "都道府県別森林率・人工林率", "林野庁", xlsx_url,
              "forest", "excel_download", "csv/jsonl",
              "公共データ利用規約(第1.0版)PDL1.0 (農林水産省ウェブサイトのコンテンツ利用規約、"
              "出典: https://www.maff.go.jp/j/use/link.html ; 出典記載必須・改変時の明記必須)",
              1, len(rows_out),
              notes=f"出典ページ: {detail_url} ; 見出し: {title_ja} ; 注記: {all_notes}")


def fetch_lidar_kanagawa_status():
    source_id = "rinya_lidar_kanagawa_status"
    url = "https://www.rinya.maff.go.jp/j/keikaku/smartforest/smart_forestry.html"
    r = get(url)
    html = r.content.decode("utf-8", errors="replace")

    rows = [{
        "source_id": source_id,
        "source_ref": url,
        "prefecture_ja": "神奈川県",
        "publication_status_ja": "公開状況不明(林野庁サイト上に都道府県別の一覧・地図は無い)",
        "data_format_ja": "",
        "access_url": "",
        "notes_ja": (
            "林野庁ページの「公開データ」節(#koukai)は原文「林野庁保有の公開データについては、"
            "G空間情報センターの林野庁ページ(外部リンク)を参照願います。また、林野庁事業等により、"
            "都道府県によるオープンデータ化も進めており、併せて、G空間情報センターの都道府県ページも"
            "ご覧ください。」とあり、実体データの所在はG空間情報センター(https://www.geospatial.jp/ckan/"
            "organization/rinya 等)のみに集約されている。G空間情報センターは本タスクでは別の並行エージェントの"
            "担当ソースであるため、本タスクからはアクセスせず深追いしていない。"
            "「申請に基づく提供データ」節(#teikyou)の問合せ先は「1.石川県能登地域」「2.その他」の2区分のみで"
            "神奈川県固有の記載なし。令和5年度委託事業のWEB-GIS実証(iframe: https://webgis-rashinban-mori.com/)"
            "対象県は原文「栃木県、兵庫県及び高知県」のみで神奈川県は対象外。"
        ),
    }]

    fieldnames = list(rows[0].keys())
    csv_path = PROC / f"{source_id}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)
    print(f"  [write] {csv_path} {len(rows)} rows")
    write_jsonl(source_id, rows)

    register(source_id, "森林情報オープンデータ化(航空レーザ計測)神奈川県公開状況", "林野庁", url,
              "forest", "html_scrape", "csv/jsonl",
              "公共データ利用規約(第1.0版)PDL1.0 (農林水産省ウェブサイトのコンテンツ利用規約、"
              "出典: https://www.maff.go.jp/j/use/link.html ; G空間情報センター側の実データは別ライセンス条件の可能性あり未確認)",
              1, len(rows),
              notes="実体(LiDAR点群・DEM等)は巨大なため未取得。目録のみ。神奈川県の公開状況は林野庁サイト上では直接確認できず、"
                    "実体はG空間情報センター経由(別エージェント担当のため本タスクでは未確認)。")


if __name__ == "__main__":
    fetch_prefecture_forest_stats()
    fetch_lidar_kanagawa_status()
